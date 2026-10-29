import io
import pytest
from bs4 import BeautifulSoup
from backend.article_dates import calendar_date, publication_from_text, publication_from_html
from backend.documents import extract_document
from backend.schemas import ArticleInput


@pytest.mark.parametrize("value,expected", [
    ("2026-10-28T08:00:00+05:30", "2026-10-28"),
    ("Oct 28, 2026 08:00 IST", "2026-10-28"),
    ("28 October 2026", "2026-10-28"),
    ("October 28th, 2026", "2026-10-28"),
    ("2026-02-30", None), ("February 30, 2026", None),
    ("28/10/26", None), (None, None),
])
def test_calendar_dates(value, expected):
    assert calendar_date(value) == expected


@pytest.mark.parametrize("text,expected", [
    ("Published OnOct 28, 2026 08:00 IST Last Updated OnOct 29, 2026", "2026-10-28"),
    ("Publication date: 2026-10-28\nThe report covers a cyclone.", "2026-10-28"),
    ("Published: 28 October 2026\nThe cyclone arrived yesterday.", "2026-10-28"),
    ("Updated 6:40 AM GMT+5:30, October 29, 2026", None),
    ("The cyclone made landfall on October 28, 2026.", None),
    ("The report was published on October 28, 2026.", None),
    ("Published: February 30, 2026", None),
])
def test_only_explicit_publication_labels(text, expected):
    assert publication_from_text(text) == expected


@pytest.mark.parametrize("html,expected", [
    ('<meta property="article:published_time" content="2026-10-28T08:00:00Z">', "2026-10-28"),
    ('<script type="application/ld+json">{"@graph":[{"@type":"NewsArticle","datePublished":"2026-10-28","dateModified":"2026-10-29"}]}</script>', "2026-10-28"),
    ('<script type="application/ld+json">{"@type":"Event","datePublished":"2026-10-28"}</script>', None),
    ('<script type="application/ld+json">invalid</script><time itemprop="datePublished" datetime="2026-10-28"></time>', "2026-10-28"),
    ('<time class="published">October 28, 2026</time>', "2026-10-28"),
    ('<meta property="article:modified_time" content="2026-10-29"><time datetime="2026-10-28">An event</time>', None),
    ('<meta property="article:published_time" content="bad"><p>Published: October 28, 2026</p>', "2026-10-28"),
])
def test_html_publication_metadata(html, expected):
    assert publication_from_html(BeautifulSoup(html, "html.parser")) == expected


def test_manual_date_takes_precedence():
    body = dict(title="News", publisher="Example", text="Published: October 28, 2026. Residents were evacuated ahead of the storm.")
    assert ArticleInput(**body).published_at == "2026-10-28"
    assert ArticleInput(**body, published_at="2026-10-27").published_at == "2026-10-27"


def test_docx_uses_article_label_not_file_creation_date():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Published OnOct 28, 2026 08:00 IST")
    doc.add_paragraph("Residents were evacuated ahead of a severe cyclone affecting the coast.")
    output = io.BytesIO()
    doc.save(output)
    result = extract_document("news.docx", output.getvalue())
    assert result["published_at"] == "2026-10-28"
    assert "Confirm" in result["warnings"][0]
