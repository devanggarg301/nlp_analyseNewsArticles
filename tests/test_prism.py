import time
from unittest.mock import Mock
import pytest
import httpx
from fastapi.testclient import TestClient
from backend import store, models, llm, fetcher
from backend.analysis import split_passages, classify, attribution, merge_entities, align_claims, coverage, enrich
from backend.demo import DEMO
from backend.app import app
from backend.retrieval import answer, is_workspace_count, retrieve, question_intent

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(store,"DB",tmp_path/"test.sqlite")
    monkeypatch.setattr(models,"encoder",lambda:None)
    monkeypatch.setattr(models,"ner",lambda:None)
    monkeypatch.setattr(models,"nli",lambda:None)
    monkeypatch.setattr(models,"bias",lambda:(None,"Lexicon fallback"))
    monkeypatch.setattr(llm,"configured",lambda:False)
    with TestClient(app) as client:
        yield client

def articles():
    return [dict(a,id=f"a{i}") for i,a in enumerate(DEMO["articles"])]

def test_offsets_preserve_original_evidence():
    a=articles()[0]
    passages=split_passages(a)
    assert len(passages)>8
    assert all(a["text"][p["start"]:p["end"]]==p["text"] for p in passages)
    assert passages[0]["text"].startswith("Hundreds of residents")
    assert "residents" in passages[0]["text"]

def test_quoted_judgement_has_speaker():
    result=attribution('Maya Chen said, “The proposal is dangerous.”')
    assert result["speaker"]=="Maya Chen"
    assert result["kind"]=="direct_quote"
    result=attribution("The minister described the protesters as dangerous.")
    assert result["kind"]=="reported_speech"
    assert result["speaker"] is None

def test_sentence_boundaries_preserve_abbreviations(monkeypatch):
    monkeypatch.setattr(models,"ner",lambda:None)
    text="Dr. Chen arrived at 10 a.m. and spoke about U.S. policy. The cost was 2.5 million."
    passages=split_passages({"id":"a","text":text})
    assert len(passages)==2
    assert passages[0]["text"].startswith("Dr. Chen")
    assert all(text[p["start"]:p["end"]]==p["text"] for p in passages)

def test_negative_event_is_not_entity_polarity(client):
    ws=client.post("/api/examples/riverton").json()
    result=run(client,ws["id"])
    assert all(p["label"]=="unclear" for e in result["analysis"]["entities"] for p in e["portrayals"])

def run(client,wid):
    response=client.post(f"/api/workspaces/{wid}/analyze",json={"confirm_unrelated":True})
    assert response.status_code==200
    jid=response.json()["job"]["id"]
    for _ in range(100):
        job=client.get("/api/jobs/"+jid).json()
        if job["status"]=="failed":
            pytest.fail(str(job))
        if job["status"]=="complete":
            return client.get("/api/workspaces/"+wid).json()
        time.sleep(.03)
    pytest.fail("Job did not complete")

def test_full_demo_and_exports(client):
    ws=client.post("/api/examples/riverton").json()
    result=run(client,ws["id"])
    assert result["analysis"]["counts"]["articles"]==3
    assert result["analysis"]["counts"]["flagged"]>0
    assert result["analysis"]["coverage"]
    assert result["analysis"]["claims"]
    markdown=client.get(f'/api/workspaces/{ws["id"]}/export?format=markdown')
    assert "Fictional teaching example" in markdown.text
    assert "Evidence index" in markdown.text
    export=client.get(f'/api/workspaces/{ws["id"]}/export?format=json').json()
    assert export["analysis"]["passages"][0]["id"] in markdown.text

def test_duplicates_and_article_limits(client):
    ws=client.post("/api/workspaces",json={"title":"Council event"}).json()
    body=DEMO["articles"][0]
    assert client.post(f'/api/workspaces/{ws["id"]}/articles',json=body).status_code==200
    assert client.post(f'/api/workspaces/{ws["id"]}/articles',json=body).status_code==409
    response=client.post(f'/api/workspaces/{ws["id"]}/articles',json=dict(body,text="word "*3001))
    assert response.status_code==422

def test_reanalysis_required_after_edit(client):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    article=ws["articles"][0]
    result=client.patch(f'/api/workspaces/{ws["id"]}/articles/{article["id"]}',json=dict(article,title="Updated headline")).json()
    assert result["analysis"] is None
    assert not result["chat"]

@pytest.mark.parametrize("url",["http://127.0.0.1/a","http://[::1]/","http://169.254.169.254/","http://10.0.0.1/","http://192.168.1.1/","http://user:password@example.com","file:///etc/passwd","https://example.com:8443/"])
def test_ssrf_rejected(url):
    with pytest.raises(ValueError):
        fetcher.public_target(url)

def test_redirect_to_private_host_is_checked(monkeypatch):
    calls=[]
    def request(url):
        calls.append(url)
        if len(calls)==1:
            return 302,{"location":"http://127.0.0.1/admin"},b""
        fetcher.public_target(url)
    monkeypatch.setattr(fetcher,"request_page",request)
    with pytest.raises(ValueError):
        fetcher.extract_url("https://publisher.example/article")
    assert len(calls)==2

def test_extraction_failure_returns_paste_fallback(client,monkeypatch):
    monkeypatch.setattr(fetcher,"request_page",lambda url:(403,{},b""))
    # request_page normally raises on 403; simulate extractor failure at API boundary.
    import backend.app as module
    monkeypatch.setattr(module,"extract_url",Mock(side_effect=ValueError("Paste the article text instead.")))
    response=client.post("/api/extract",json={"url":"https://example.com/story"})
    assert response.status_code==422
    assert "Paste" in response.json()["detail"]

def test_insufficient_evidence(client):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    response=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"What is the orbital radius of Jupiter?"})
    assert response.status_code==200
    assert response.json()["mode"]=="insufficient_evidence"

def test_source_balanced_retrieval_and_citations(client):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    result=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"What do the articles report about bus fares and council?"}).json()
    assert result["mode"]=="extractive"
    assert len(set(e["article_id"] for e in result["evidence"]))==3
    valid={p["id"] for p in ws["analysis"]["passages"]}
    assert all(pid in valid for c in result["claims"] for pid in c["evidence"])

def test_api_failure_keeps_analysis(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    monkeypatch.setattr(llm,"configured",lambda:True)
    monkeypatch.setattr(llm,"complete_json",Mock(side_effect=RuntimeError("upstream failure")))
    response=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"What are the bus fare plans?"})
    assert response.status_code==200
    assert response.json()["mode"]=="extractive_fallback"
    assert client.get(f'/api/workspaces/{ws["id"]}').json()["analysis"]

def test_invalid_llm_citations_are_not_shown(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    monkeypatch.setattr(llm,"configured",lambda:True)
    monkeypatch.setattr(llm,"complete_json",lambda *args:{"claims":[{"text":"An unsupported answer","evidence":["made-up"]}]})
    response=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"What are the bus fare plans?"}).json()
    assert response["mode"]=="extractive_fallback"
    assert not any(c["text"]=="An unsupported answer" for c in response["claims"])

def test_different_temporal_context_not_reported_as_contradiction(monkeypatch):
    monkeypatch.setattr(models,"encoder",lambda:None)
    monkeypatch.setattr(models,"nli",lambda:lambda *a,**k:[{"label":"contradiction","score":.99},{"label":"contradiction","score":.99}])
    p=[{"id":"a:p1","article_id":"a","text":"The meeting was held on Monday."},{"id":"b:p1","article_id":"b","text":"The meeting was held on Tuesday."}]
    result=align_claims(p)
    assert result
    assert result[0]["label"]=="unresolved"
    assert result[0]["context_warning"]

def test_merge_invalidates_target_specific_judgements():
    values=[{"name":n,"aliases":[n],"evidence":["a:p1"],"portrayals":[{"article_id":"a","label":label,"evidence":["a:p1"]}]} for n,label in [("Maya Chen","favourable"),("Chen","unfavourable")]]
    result=merge_entities(values,["Maya Chen","Chen"],"Maya Chen")
    assert len(result)==1
    assert result[0]["portrayals"][0]["label"]=="unclear"

def test_upload_csv_and_reject_invalid(client):
    response=client.post("/api/preview-upload",files={"file":("articles.csv",'title,publisher,text\nHeadline,Source,"An article with enough text to add to the workspace."',"text/csv")})
    assert response.status_code==200
    assert response.json()["articles"][0]["publisher"]=="Source"
    assert client.post("/api/preview-upload",files={"file":("a.pdf",b"pdf","application/pdf")}).status_code==422

def test_rag_success_has_valid_evidence(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    monkeypatch.setattr(llm,"configured",lambda:True)
    def generated(system,payload):
        evidence=payload["evidence"][0]
        return {"insufficient_evidence":False,"claims":[{"text":evidence["text"],"evidence":[evidence["id"]]}]}
    monkeypatch.setattr(llm,"complete_json",generated)
    result=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"What are the bus fare plans?"}).json()
    assert result["mode"]=="rag"
    valid={p["id"] for p in ws["analysis"]["passages"]}
    assert all(pid in valid for c in result["claims"] for pid in c["evidence"])

def test_ner_failure_is_recoverable(client,monkeypatch):
    monkeypatch.setattr(models,"ner",Mock(side_effect=RuntimeError("missing weights")))
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    assert ws["analysis"]["passages"]
    assert any("Entity extraction failed" in w for w in ws["analysis"]["warnings"])

def test_five_article_cap_and_minimum(client):
    ws=client.post("/api/workspaces",json={"title":"A sample event"}).json()
    assert client.post(f'/api/workspaces/{ws["id"]}/analyze',json={}).status_code==422
    for i in range(5):
        response=client.post(f'/api/workspaces/{ws["id"]}/articles',json=dict(DEMO["articles"][0],text=DEMO["articles"][0]["text"]+f" Unique item {i}."))
        assert response.status_code==200
    response=client.post(f'/api/workspaces/{ws["id"]}/articles',json=dict(DEMO["articles"][0],text=DEMO["articles"][0]["text"]+" Final extra."))
    assert response.status_code==422

def test_workspace_count_question(client):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    result=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"How many sources do we have?"}).json()
    assert result["mode"]=="workspace_metadata"
    assert "3 articles" in result["claims"][0]["text"]

@pytest.mark.parametrize("question",[
    "how many documents do we have?",
    "How many docs are there?",
    "How many files have we uploaded?",
    "What is the total number of documents?",
    "Number of documents in this workspace",
    "Total document count",
    "How many articles are selected?",
    "How many publishers do we have?",
])
def test_document_count_wording(question):
    analysis={"passages":[]}
    result=answer(question,analysis,articles())
    assert result["mode"]=="workspace_metadata"
    assert "3 articles" in result["claims"][0]["text"]

@pytest.mark.parametrize("question",[
    "How many documents mention police?",
    "How many articles discuss traffic disruption?",
    "How many documents are there about the council?",
])
def test_qualified_count_is_not_confused_with_workspace_size(question):
    assert not is_workspace_count(question)

def test_current_config_does_not_show_historical_missing_key_warning(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    assert any(note.startswith("No API key configured:") for note in ws["analysis"]["warnings"])
    monkeypatch.setattr(llm,"configured",lambda:True)
    response=client.get(f'/api/workspaces/{ws["id"]}').json()
    assert not any(note.startswith("No API key configured:") for note in response["analysis"]["warnings"])
    assert any("created before API setup" in note for note in response["analysis"]["warnings"])
    # The display correction must not pretend the original analysis used an API.
    original=store.get(ws["id"])
    assert original["analysis"]["llm_configured_at_analysis"] is False
    assert any(note.startswith("No API key configured:") for note in original["analysis"]["warnings"])

def test_reanalysis_after_api_setup_preserves_chat(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"How many documents do we have?"})
    before=store.get(ws["id"])["chat"]
    monkeypatch.setattr(llm,"configured",lambda:True)
    monkeypatch.setattr(llm,"complete_json",lambda *args:{"portrayals":[],"phrases":[]})
    after=run(client,ws["id"])
    assert after["chat"]==before
    assert after["analysis"]["llm_configured_at_analysis"] is True
    assert not any("No API key configured" in note for note in after["analysis"]["warnings"])

@pytest.mark.parametrize("shape",["bare_array","nested_rows","wrapped_array"])
def test_entity_analysis_accepts_provider_array_responses(monkeypatch,shape):
    entity={"name":"Maya Chen","aliases":["Maya Chen"],"evidence":["a:p1"],"portrayals":[]}
    passage={"id":"a:p1","article_id":"a","text":"Maya Chen spoke at the meeting."}
    row={"entity":"Maya Chen","article_id":"a","label":"neutral","evidence":["a:p1"],"explanation":"A neutral mention."}
    response={"bare_array":[row],"nested_rows":{"portrayals":[[row]]},"wrapped_array":[{"portrayals":[row]}]}[shape]
    monkeypatch.setattr(llm,"complete_json",lambda *args:response)
    enrich([entity],[passage],[],[{"id":"a"}])
    assert entity["portrayals"][0]["label"]=="neutral"
    assert entity["portrayals"][0]["evidence"]==["a:p1"]

def test_invalid_entity_evidence_does_not_count_as_success(monkeypatch):
    entity={"name":"Maya Chen","aliases":[],"evidence":["a:p1"],"portrayals":[]}
    monkeypatch.setattr(llm,"complete_json",lambda *args:[{"entity":"Maya Chen","article_id":"a","label":"favourable","evidence":["invented"]}])
    with pytest.raises(ValueError,match="No valid"):
        enrich([entity],[{"id":"a:p1","article_id":"a","text":"Maya Chen spoke."}],[],[{"id":"a"}])
    assert not entity["portrayals"]

def test_rag_accepts_bare_claim_array(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    monkeypatch.setattr(llm,"configured",lambda:True)
    monkeypatch.setattr(llm,"complete_json",lambda system,payload:[{"text":payload["evidence"][0]["text"],"evidence":[payload["evidence"][0]["id"]]}])
    result=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"What are the bus fare plans?"}).json()
    assert result["mode"]=="rag"
    assert result["claims"][0]["evidence"]

def test_transient_provider_error_is_retried(monkeypatch):
    monkeypatch.setattr(llm,"configured",lambda:True)
    monkeypatch.setenv("LLM_API_KEY","test-token")
    monkeypatch.setenv("LLM_MODEL","test-model")
    monkeypatch.setattr(llm.time,"sleep",lambda seconds:None)
    request=httpx.Request("POST","https://example.invalid/chat/completions")
    responses=[httpx.Response(503,request=request),httpx.Response(200,request=request,json={"choices":[{"message":{"content":'{"ok":true}'}}]})]
    post=Mock(side_effect=responses)
    monkeypatch.setattr(llm.httpx,"post",post)
    assert llm.complete_json("Return JSON",{})=={"ok":True}
    assert post.call_count==2

def test_authentication_error_is_not_retried(monkeypatch):
    monkeypatch.setattr(llm,"configured",lambda:True)
    monkeypatch.setenv("LLM_API_KEY","test-token")
    monkeypatch.setenv("LLM_MODEL","test-model")
    request=httpx.Request("POST","https://example.invalid/chat/completions")
    post=Mock(return_value=httpx.Response(401,request=request))
    monkeypatch.setattr(llm.httpx,"post",post)
    with pytest.raises(httpx.HTTPStatusError):
        llm.complete_json("Return JSON",{})
    assert post.call_count==1

@pytest.mark.parametrize("question",["whats the main topic of discussion","what is the topic","what is article about"])
def test_overview_questions_use_corpus_not_similarity_cutoff(client,question):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    result=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":question}).json()
    assert result["mode"]!="insufficient_evidence"
    assert len(result["evidence"])==len(ws["analysis"]["passages"])

def test_affected_place_question_has_full_context(client,monkeypatch):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    monkeypatch.setattr(llm,"configured",lambda:True)
    def generated(system,payload):
        assert payload["intent"]=="affected_places"
        assert "geographical locations" in system
        assert len(payload["evidence"])==len(ws["analysis"]["passages"])
        p=payload["evidence"][0]
        return {"claims":[{"text":"Riverton City Hall is the location described.","evidence":[p["id"]]}]}
    monkeypatch.setattr(llm,"complete_json",generated)
    result=client.post(f'/api/workspaces/{ws["id"]}/chat',json={"question":"which place is affected"}).json()
    assert result["mode"]=="rag"

def test_docx_upload_preview(client):
    import io
    from docx import Document
    doc=Document()
    doc.core_properties.title="Cyclone news report"
    doc.add_paragraph("Officials reported heavy rainfall in the coastal districts. Residents were moved to temporary shelters before landfall.")
    stream=io.BytesIO();doc.save(stream)
    result=client.post("/api/preview-upload",files={"file":("article.docx",stream.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert result.status_code==200
    assert "temporary shelters" in result.json()["articles"][0]["text"]
    assert result.json()["articles"][0]["title"]=="Cyclone news report"

def test_pdf_upload_and_scan_message(client):
    import io
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject,NameObject,DictionaryObject
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    font=DictionaryObject({NameObject("/Type"):NameObject("/Font"),NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
    page[NameObject("/Resources")]=DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b"BT /F1 12 Tf 50 700 Td (Heavy rainfall affected coastal villages. Officials opened shelters and evacuated residents before the cyclone reached the coast.) Tj ET")
    page[NameObject("/Contents")]=writer._add_object(stream)
    buf=io.BytesIO();writer.write(buf)
    result=client.post("/api/preview-upload",files={"file":("news.pdf",buf.getvalue(),"application/pdf")})
    assert result.status_code==200
    assert "coastal villages" in result.json()["articles"][0]["text"]
    blank=PdfWriter();blank.add_blank_page(width=612,height=792);buf=io.BytesIO();blank.write(buf)
    response=client.post("/api/preview-upload",files={"file":("scan.pdf",buf.getvalue(),"application/pdf")})
    assert response.status_code==422
    assert "OCR" in response.json()["detail"]

def test_legacy_doc_instructions(client):
    response=client.post("/api/preview-upload",files={"file":("news.doc",b"old word file","application/msword")})
    assert response.status_code==422
    assert ".docx" in response.json()["detail"]

def test_portable_model_manifest_when_neural_weights_missing(tmp_path,monkeypatch):
    import json
    monkeypatch.setattr(models,"MODELS",tmp_path)
    (tmp_path/"selected.json").write_text(json.dumps({"kind":"transformer","path":"missing-roberta","name":"RoBERTa"}))
    (tmp_path/"default-selected.json").write_text(json.dumps({"kind":"sklearn","path":"baseline.joblib","name":"Baseline"}))
    (tmp_path/"baseline.joblib").write_bytes(b"fixture")
    assert models.selection_manifest().name=="default-selected.json"

def test_manual_entity_grouping_survives_reanalysis(client):
    ws=run(client,client.post("/api/examples/riverton").json()["id"])
    names=[e["name"] for e in ws["analysis"]["entities"][:2]]
    response=client.post(f'/api/workspaces/{ws["id"]}/entities/merge',json={"names":names,"canonical":names[0]})
    assert response.status_code==200
    assert response.json()["entity_merges"][0]["names"]==names
    after=run(client,ws["id"])
    assert names[1] not in [e["name"] for e in after["analysis"]["entities"]]

def test_exact_death_count_disagreement_in_complete_analysis(client):
    ws=client.post("/api/workspaces",json={"title":"Same-event death count test"}).json()
    for publisher,text in [("Article A","Exactly five people had died by 6 p.m."),("Article B","Exactly three people had died by 6 p.m.")]:
        result=client.post(f'/api/workspaces/{ws["id"]}/articles',json={"title":"Reported death toll","publisher":publisher,"text":text})
        assert result.status_code==200
    ws=run(client,ws["id"])
    assert ws["analysis"]["counts"]["flagged"]==0
    conflicts=[c for c in ws["analysis"]["claims"] if c["label"]=="potential_contradiction"]
    assert len(conflicts)==1
    assert conflicts[0]["method"]=="Numerical count and scope comparison"


